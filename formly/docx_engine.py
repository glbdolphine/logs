"""
formly/docx_engine.py
============================
Everything that touches the .docx template: mixed Bangla/English font
handling, finding a label in a paragraph/cell and inserting a value after
it, the header line (সুত্র নং / তারিখ), the site-visit sentence, the cable
quantity row, and MAC address cleanup.

This is the same engine the original app.py used — logic unchanged, just
pulled out into its own module and reading its label text from config.py
instead of hardcoding it here.
"""

import re
import unicodedata
from datetime import datetime

from config import LABELS, HEADER_SOURCE_NO_LABELS, HEADER_DATE_LABELS, BANGLA_FONT, ENGLISH_FONT


def nfc(s):
    return unicodedata.normalize('NFC', s) if s else s


def sanitize_text(s):
    if not s:
        return s
    return s.encode('utf-8', 'ignore').decode('utf-8')


def sanitize_filename(name):
    name = re.sub(r'[\\/:*?"<>|]+', '', name)
    name = re.sub(r'\s+', '_', name).strip('_')
    return name[:60] if name else "form"


# ---------------------------------------------------------------------------
# Mixed Bangla/English font handling
# ---------------------------------------------------------------------------

def _is_bangla_char(ch):
    return '\u0980' <= ch <= '\u09FF'


def _char_script(ch):
    if _is_bangla_char(ch):
        return True
    if ch.isalpha() or ch.isdigit():
        return False
    return None


def _split_by_script(text):
    segments = []
    current = ''
    current_is_bangla = None
    for ch in text:
        is_bangla = _char_script(ch)
        if is_bangla is None:
            is_bangla = current_is_bangla if current_is_bangla is not None else False
        if current_is_bangla is None:
            current_is_bangla = is_bangla
            current = ch
        elif is_bangla == current_is_bangla:
            current += ch
        else:
            segments.append((current, current_is_bangla))
            current = ch
            current_is_bangla = is_bangla
    if current:
        segments.append((current, current_is_bangla))
    return segments if segments else [(text, False)]


def _set_run_font(run, font_name):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(attr), font_name)


def _set_run_size(run, size):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    run.font.size = size
    rPr = run._element.get_or_add_rPr()
    half_points = str(int(round(size.pt * 2)))
    szCs = rPr.find(qn('w:szCs'))
    if szCs is None:
        szCs = OxmlElement('w:szCs')
        rPr.append(szCs)
    szCs.set(qn('w:val'), half_points)


def _copy_run_format(src_run, dst_run):
    from docx.shared import Pt
    dst_run.bold = src_run.bold
    dst_run.italic = src_run.italic
    dst_run.underline = src_run.underline
    _set_run_size(dst_run, src_run.font.size or Pt(11))
    try:
        if src_run.font.color and src_run.font.color.rgb is not None:
            dst_run.font.color.rgb = src_run.font.color.rgb
    except Exception:
        pass


def apply_mixed_fonts(paragraph, run):
    text = run.text
    if not text:
        return
    segments = _split_by_script(text)
    first_text, first_is_bangla = segments[0]
    run.text = first_text
    _set_run_font(run, BANGLA_FONT if first_is_bangla else ENGLISH_FONT)
    prev_element = run._element
    for seg_text, seg_is_bangla in segments[1:]:
        new_run = paragraph.add_run(seg_text)
        _copy_run_format(run, new_run)
        _set_run_font(new_run, BANGLA_FONT if seg_is_bangla else ENGLISH_FONT)
        prev_element.addnext(new_run._element)
        prev_element = new_run._element


def _capture_format(paragraph):
    from docx.shared import Pt
    if paragraph.runs:
        r = paragraph.runs[0]
        color = None
        try:
            if r.font.color and r.font.color.type is not None:
                color = r.font.color.rgb
        except Exception:
            pass
        return {'size': r.font.size or Pt(11), 'bold': r.font.bold,
                'italic': r.font.italic, 'underline': r.font.underline, 'color': color}
    return {'size': Pt(11), 'bold': None, 'italic': None, 'underline': None, 'color': None}


def _apply_captured_format(run, fmt):
    _set_run_size(run, fmt['size'])
    run.font.bold = fmt['bold']
    run.font.italic = fmt['italic']
    run.font.underline = fmt['underline']
    if fmt['color'] is not None:
        try:
            run.font.color.rgb = fmt['color']
        except Exception:
            pass


def set_paragraph_text_preserving_format(paragraph, new_text):
    new_text = sanitize_text(new_text)
    fmt = _capture_format(paragraph)
    paragraph.text = new_text
    if not paragraph.runs:
        return
    run = paragraph.runs[0]
    _apply_captured_format(run, fmt)
    apply_mixed_fonts(paragraph, run)


# ---------------------------------------------------------------------------
# Bangla dates
# ---------------------------------------------------------------------------

def to_bangla_digits(number):
    bangla_digits = {'0': '০', '1': '১', '2': '২', '3': '৩', '4': '৪',
                      '5': '৫', '6': '৬', '7': '৭', '8': '৮', '9': '৯'}
    return ''.join(bangla_digits.get(ch, ch) for ch in str(number))


def get_bangla_date(dt):
    day = to_bangla_digits(f"{dt.day:02d}")
    month = to_bangla_digits(f"{dt.month:02d}")
    year = to_bangla_digits(dt.year)
    return f"{day}/{month}/{year}"


def get_today_bangla_date():
    return get_bangla_date(datetime.now())


# ---------------------------------------------------------------------------
# MAC address cleanup
# ---------------------------------------------------------------------------

_MAC_CANDIDATE_RE = re.compile(
    r'(?<![0-9A-Fa-f:\-])([0-9A-Fa-f](?:[:\-]?[0-9A-Fa-f]){11})(?![0-9A-Fa-f])'
)


def _format_mac(hex_only):
    pairs = [hex_only[i:i + 2] for i in range(0, 12, 2)]
    return ':'.join(pairs).upper()


def normalize_mac(raw):
    """Best-effort MAC cleanup, used on whatever comes out of the sheet or
    manual entry — accepts with/without separators, mixed case."""
    if not raw:
        return ""
    candidate = _MAC_CANDIDATE_RE.search(raw)
    if candidate:
        hex_only = re.sub(r'[^0-9A-Fa-f]', '', candidate.group(1))
        if len(hex_only) == 12:
            return _format_mac(hex_only)
    return raw.strip()


# ---------------------------------------------------------------------------
# Label finding / value insertion
# ---------------------------------------------------------------------------

def _find_label(text, label_variants):
    ntext = nfc(text)
    for label in label_variants:
        nlabel = nfc(label)
        idx = ntext.find(nlabel)
        if idx != -1:
            return nlabel, idx
    return None, -1


def _build_insertion(label, value):
    if label and label[-1] in ':।ঃ':
        return f" {value}"
    return f": {value}"


def _insert_value(text, label_variants, value, stop_labels=None):
    ntext = nfc(text)
    label, idx = _find_label(ntext, label_variants)
    if label is None:
        return None
    end = idx + len(label)
    remainder = ntext[end:]

    stop_idx = len(remainder)
    found_stop = False
    if stop_labels:
        for stop in stop_labels:
            pos = remainder.find(nfc(stop))
            if pos != -1 and pos < stop_idx:
                stop_idx = pos
                found_stop = True

    insertion = _build_insertion(label, value)
    tail = remainder[stop_idx:]

    if found_stop:
        original_gap_width = stop_idx
        pad = max(original_gap_width - len(insertion), 1)
        insertion = insertion + (" " * pad)
    elif tail and not tail[0].isspace():
        tail = " " + tail

    return ntext[:end] + insertion + tail


def update_text_in_paragraphs(doc, label_variants, new_value):
    for paragraph in doc.paragraphs:
        new_text = _insert_value(paragraph.text, label_variants, new_value)
        if new_text is not None:
            set_paragraph_text_preserving_format(paragraph, new_text)
            return True
    return False


def update_cell_or_paragraph(doc, label_variants, new_value):
    if not new_value:
        return False
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    new_text = _insert_value(paragraph.text, label_variants, new_value)
                    if new_text is not None:
                        set_paragraph_text_preserving_format(paragraph, new_text)
                        return True
    if update_text_in_paragraphs(doc, label_variants, new_value):
        return True
    return False


def fill_header_line(doc, source_no, header_date):
    for paragraph in doc.paragraphs:
        ntext = nfc(paragraph.text)
        if nfc("সুত্র") in ntext and nfc("তারিখ") in ntext:
            text = ntext
            if source_no:
                updated = _insert_value(text, HEADER_SOURCE_NO_LABELS, source_no, stop_labels=HEADER_DATE_LABELS)
                if updated is not None:
                    text = updated
            if header_date:
                updated = _insert_value(text, HEADER_DATE_LABELS, header_date)
                if updated is not None:
                    text = updated
            if text != ntext:
                set_paragraph_text_preserving_format(paragraph, text)
            return True
    return False


def insert_name_in_site_visit_paragraph(doc, institution_name):
    if not institution_name:
        return False
    marker = nfc("এর উপরোক্ত")
    anchor = nfc("সরজমিনে পরিদর্শন")
    replacement = nfc(f"এর {institution_name} উপরোক্ত")
    for paragraph in doc.paragraphs:
        para_text_nfc = nfc(paragraph.text)
        if marker in para_text_nfc and anchor in para_text_nfc:
            new_text = para_text_nfc.replace(marker, replacement, 1)
            set_paragraph_text_preserving_format(paragraph, new_text)
            return True
    return False


def set_cable_quantity(doc, qty):
    if not qty:
        return False
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 4:
                continue
            if "Optical Fiber Cable" in cells[1].text:
                qty_paragraph = cells[3].paragraphs[0]
                set_paragraph_text_preserving_format(qty_paragraph, str(qty))
                return True
    return False


# ---------------------------------------------------------------------------
# Top-level: fill one form
# ---------------------------------------------------------------------------

def fill_form(template_path, output_path, data):
    from docx import Document
    doc = Document(template_path)
    full_address = data.get('name', '')

    fill_map = [
        (LABELS['edc_book'], data.get('source_no', '')),
        (LABELS['connection_nms'], data.get('connection_nms', '')),
        (LABELS['address'], full_address),
        (LABELS['lat_long'], data.get('lat_long', '')),
        (LABELS['union'], data.get('union', '')),
        (LABELS['upazila'], data.get('upazila', '')),
        (LABELS['institution_code'], data.get('institution_code', '')),
        (LABELS['date'], data.get('date', '')),
        (LABELS['router_mac'], data.get('router_mac', '')),
        (LABELS['router_serial'], data.get('router_serial', '')),
    ]
    for label_variants, value in fill_map:
        if value:
            update_cell_or_paragraph(doc, label_variants, value)

    fill_header_line(doc, data.get('source_no', ''), data.get('date', ''))
    insert_name_in_site_visit_paragraph(doc, data.get('name', ''))
    set_cable_quantity(doc, data.get('cable_qty', ''))

    cable_brand = data.get('cable_brand', 'BRB')
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "Optical Fiber Cable Brand:" in paragraph.text:
                        new_text = _insert_value(paragraph.text, ["Optical Fiber Cable Brand:"], cable_brand)
                        if new_text:
                            set_paragraph_text_preserving_format(paragraph, new_text)
                        break

    doc.save(output_path)
